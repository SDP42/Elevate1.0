"""Read the organiser's XLSX and prepare private normalized JSON. No DB writes.
Requires openpyxl. Usage: python3 db/prepare_rsvp.py input.xlsx output-directory
"""
import hashlib
import json
import os
import re
import sys
from collections import Counter
from datetime import datetime, timezone, timedelta
from pathlib import Path
from openpyxl import load_workbook

ACRONYMS = {"ai": "AI", "ml": "ML", "aiml": "AI/ML", "ds": "DS", "aids": "AI/DS", "it": "IT", "cse": "CSE", "cs": "CS", "sies": "SIES", "vjti": "VJTI", "djsce": "DJSCE", "apsit": "APSIT", "spit": "SPIT", "mpstme": "MPSTME", "svkm": "SVKM", "mmcc": "MMCC", "lr": "LR", "vesit": "VESIT"}

def clean(value):
    if value is None: return ""
    if isinstance(value, float) and value.is_integer(): value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()

def key(value): return re.sub(r"[^a-z0-9]", "", clean(value).lower())

def title(value):
    text = clean(value).title()
    text = re.sub(r"\b(\d+)(St|Nd|Rd|Th)\b", lambda m: m.group(1) + m.group(2).lower(), text)
    return re.sub(r"\b[A-Za-z]+\b", lambda m: ACRONYMS.get(m.group().lower(), m.group()), text)

def college(value):
    text = re.sub(r"\+?91[\s-]*\d{5}[\s-]*\d{5}", "", clean(value)).strip()
    ident = key(text)
    if ident in ("", "na", "none"): return None
    if ident == "djsce" or "sanghvi" in ident:
        return "Dwarkadas J. Sanghvi College of Engineering"
    if ident == "vjti" or "veermatajijabai" in ident:
        return "Veermata Jijabai Technological Institute (VJTI)"
    if "vidyavardhini" in ident: return "Vidyavardhini's College of Engineering and Technology"
    if "shah" in ident and "anchor" in ident: return "Shah & Anchor Kutchhi Engineering College"
    if "thadomal" in ident: return "Thadomal Shahani Engineering College"
    if "thakurcollege" in ident: return "Thakur College of Engineering and Technology"
    if "siesgraduateschool" in ident: return "SIES Graduate School of Technology"
    if "stfrancis" in ident: return "St. Francis Institute of Technology"
    if "vidyalankarinstitute" in ident: return "Vidyalankar Institute of Technology"
    if "lrtiwari" in ident: return "Shree LR Tiwari College of Engineering"
    return title(text)

def phone(value):
    text = clean(value)
    digits = re.sub(r"\D", "", text)
    if len(digits) == 11 and digits.startswith("0"): digits = digits[1:]
    if len(digits) == 10: return "+91" + digits
    if len(digits) == 12 and digits.startswith("91"): return "+" + digits
    return text or None

def private_json(path, data):
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "w") as output: json.dump(data, output, ensure_ascii=False, indent=2, default=str); output.write("\n")
    os.chmod(path, 0o600)

def prepare(input_path):
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook["Form responses 1"]
    iterator = sheet.iter_rows(values_only=True)
    headers = list(next(iterator)); indexes = {clean(header): i for i, header in enumerate(headers)}
    def cell(row, name): return row[indexes[name]] if indexes[name] < len(row) else None
    source_hash = hashlib.sha256(Path(input_path).read_bytes()).hexdigest()
    teams, seen_teams, seen_emails = [], set(), set()
    for row_number, row in enumerate(iterator, 2):
        if not any(value is not None for value in row): continue
        size = int(cell(row, "Team size")); team_name = clean(cell(row, "Team name"))
        issues, members = [], []
        if key(team_name) in seen_teams: issues.append("Duplicate team name")
        seen_teams.add(key(team_name))
        name_headers = ["Full name", "Full Name", "Full Name 2", "Full Name 3"]
        food_headers = ["Member 1", "Member 2", "Member 3 (if any)", "Member 4 (if any)"]
        for position in range(4):
            suffix = "" if position == 0 else f" {position + 1}"
            name = clean(cell(row, name_headers[position])).title()
            if not name: continue
            email = clean(cell(row, "Email Address" + suffix)).lower()
            food = title(cell(row, food_headers[position])) or None
            if email and email in seen_emails: issues.append(f"Member {position+1}: email appears more than once")
            if email: seen_emails.add(email)
            if position < size and food not in ("Veg", "Jain"): issues.append(f"Member {position+1}: missing or unknown food preference")
            member = {"position": position+1, "name": name, "isLead": position == 0, "email": email or None,
                      "phone": phone(cell(row, "Phone Number" + suffix)), "college": college(cell(row, "College Name" + suffix)),
                      "yearBranch": title(cell(row, "Year and Branch" + suffix)) or None, "foodPreference": food}
            members.append(member)
        if len(members) != size: issues.append(f"Declared size {size}; {len(members)} names supplied")
        stamp = cell(row, "Timestamp")
        if isinstance(stamp, datetime): stamp = stamp.replace(tzinfo=timezone(timedelta(hours=5, minutes=30))).isoformat()
        payers = list(dict.fromkeys(clean(cell(row, field)).title() for field in ["Name of the person who made the payment", "Name of the person who made the payment 2"] if clean(cell(row, field))))
        teams.append({"sourceRow": row_number, "teamName": team_name, "declaredSize": size, "respondedAt": stamp,
                      "respondentEmail": clean(cell(row, "Email address")).lower(), "attending": clean(cell(row, "Will your team attend Elevate 1.0 in person at DJSCE on 10-11 October 2026?")).startswith("Yes"),
                      "paymentPayers": payers, "termsConfirmation": clean(cell(row, "I/We agree to the following terms and conditions")),
                      "declaration": clean(cell(row, "Declaration")), "members": members, "issues": issues})
    result = {"version": 1, "sourceHash": source_hash, "sourceSheet": sheet.title, "teams": teams}
    if "drive.google.com" in json.dumps(result).lower(): raise ValueError("A Drive link remains in a non-upload field; review the source.")
    return result

if __name__ == "__main__":
    source, destination = sys.argv[1:]
    output = Path(destination); output.mkdir(parents=True, exist_ok=True); os.chmod(output, 0o700)
    data = prepare(source)
    private_json(output / "normalized.json", data)
    mapping_path = output / "mapping.json"
    if not mapping_path.exists(): private_json(mapping_path, [{"sourceTeam": team["teamName"], "username": None, "memberPositions": None} for team in data["teams"]])
    preferences = Counter(member["foodPreference"] for team in data["teams"] for member in team["members"] if member["foodPreference"])
    print(json.dumps({"teams": len(data["teams"]), "declaredParticipants": sum(team["declaredSize"] for team in data["teams"]), "submittedNames": sum(len(team["members"]) for team in data["teams"]), "foodPreferences": dict(preferences), "reviewRequired": [{"team": team["teamName"], "row": team["sourceRow"], "issues": team["issues"]} for team in data["teams"] if team["issues"]]}, indent=2))
